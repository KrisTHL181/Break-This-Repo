"""Solve problem 3 until all internal moisture values are below 0.15."""
from pathlib import Path
import sys
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _model_core import variable_material_23, integrate

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "result"


def write_book(path, times, moisture):
    wb = Workbook()
    ws = wb.active
    ws.title = "水分浓度"
    ws.cell(1, 1, "时间(s)\\到药材中心的距离(cm)")
    for j, distance in enumerate(np.arange(0, 2.01, 0.1), 2):
        ws.cell(1, j, round(float(distance), 1))
    for i, (time, row) in enumerate(zip(times, moisture), 2):
        ws.cell(i, 1, int(round(float(time))))
        for j, value in enumerate(row, 2):
            ws.cell(i, j, round(float(value), 4))
    ws.freeze_panes = "B2"
    ws.cell(1, 1).font = Font(bold=True)
    for cell in ws[1]:
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 18
    for col in range(2, 23):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 11
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def stop(C, t):
    return float(np.max(C)) <= 0.15


def main():
    # The stated 2--3 day range is extended only until the threshold is reached.
    times, _, C, max_iter = integrate(90 * 3600.0, 60.0, 12.0, variable_material_23(), stop_when=stop)
    write_book(OUT / "result3.xlsx", times, C)
    np.savez(OUT / "result3.npz", times=times, moisture=C)
    print(f"problem3: end={times[-1]/3600:.2f} h, max C={np.max(C[-1]):.6f}, max Picard iterations={max_iter}")


if __name__ == "__main__":
    main()
