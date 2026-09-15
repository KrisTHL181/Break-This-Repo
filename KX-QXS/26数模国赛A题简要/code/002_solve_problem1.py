"""Solve problem 1 and generate result/result1.xlsx."""
from pathlib import Path
import sys
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _model_core import constant_material, integrate


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "result"


def write_book(path, times, temperature, moisture):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_t = wb.active
    ws_t.title = "温度"
    ws_c = wb.create_sheet("水分浓度")
    for ws, data in ((ws_t, temperature), (ws_c, moisture)):
        ws.cell(1, 1, "时间\\到药材中心的距离(cm)")
        for j, distance in enumerate(np.arange(0, 2.01, 0.1), 2):
            ws.cell(1, j, round(float(distance), 1))
        for i, (time, row) in enumerate(zip(times, data), 2):
            ws.cell(i, 1, int(round(float(time))))
            for j, value in enumerate(row, 2):
                ws.cell(i, j, round(float(value), 4))
        ws.freeze_panes = "B2"
        ws.cell(1, 1).font = Font(bold=True)
        for cell in ws[1]:
            cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 18
        for col in range(2, 23):
            ws.column_dimensions[ws.cell(1, col).column_letter].width = 11
    wb.save(path)


def main():
    times, T, C, max_iter = integrate(1800.0, 1.0, 0.2, constant_material())
    write_book(OUT / "result1.xlsx", times, T, C)
    np.savez(OUT / "result1.npz", times=times, temperature=T, moisture=C)
    print(f"problem1: {len(times)} outputs, max Picard iterations={max_iter}")


if __name__ == "__main__":
    main()
