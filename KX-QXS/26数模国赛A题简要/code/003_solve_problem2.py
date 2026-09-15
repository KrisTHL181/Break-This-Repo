"""Solve problem 2 and generate result/result2.xlsx."""
from pathlib import Path
import sys
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _model_core import variable_material_23, integrate

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "result"


def write_sheet(ws, times, data):
    ws.cell(1, 1, "时间(s)\\到药材中心的距离(cm)")
    for j, distance in enumerate(np.arange(0, 2.01, 0.1), 2):
        ws.cell(1, j, round(float(distance), 1))
    for i, (time, row) in enumerate(zip(times, data), 2):
        ws.cell(i, 1, int(round(float(time))))
        for j, value in enumerate(row, 2):
            ws.cell(i, j, round(float(value), 4))
    ws.freeze_panes = "B2"
    ws.cell(1, 1).font = Font(bold=True)
    for cell in ws[1]:
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 20
    for col in range(2, 23):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 11


def main():
    times, T, C, max_iter = integrate(3 * 3600.0, 1.0, 0.2, variable_material_23())
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    write_sheet(wb.active, times, T)
    wb.active.title = "温度"
    ws = wb.create_sheet("水分浓度")
    write_sheet(ws, times, C)
    wb.save(OUT / "result2.xlsx")
    np.savez(OUT / "result2.npz", times=times, temperature=T, moisture=C)
    print(f"problem2: {len(times)} outputs, max Picard iterations={max_iter}")


if __name__ == "__main__":
    main()
